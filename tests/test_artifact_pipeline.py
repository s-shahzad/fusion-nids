from __future__ import annotations

import json
import sqlite3
import zipfile
from pathlib import Path

import pytest
from pypdf import PdfWriter

import src.NIDS.artifacts.intake as intake_module
from src.NIDS.artifacts.intake import ArtifactScanSummary, run_artifact_scan, run_artifact_watch
from src.NIDS.artifacts.parsers.csv_parser import parse_csv
from src.NIDS.artifacts.parsers.docx_parser import parse_docx
from src.NIDS.artifacts.parsers.exe_parser import parse_exe
from src.NIDS.artifacts.parsers.html_parser import parse_html
from src.NIDS.artifacts.parsers.json_parser import parse_json
from src.NIDS.artifacts.parsers.pdf_parser import parse_pdf
from src.NIDS.artifacts.parsers.py_parser import parse_python
from src.NIDS.artifacts.parsers.xlsx_parser import parse_xlsx
from src.NIDS.artifacts.parsers.zip_parser import parse_zip
from src.NIDS.artifacts.report import generate_artifact_report
from src.NIDS.artifacts.storage import ArtifactStore


def _make_pdf(path: Path) -> Path:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": "Fixture PDF", "/Author": "pytest"})
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def _make_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.core_properties.author = "pytest"
    document.core_properties.title = "Fixture DOCX"
    document.add_paragraph("This is a deterministic DOCX fixture.")
    document.save(str(path))
    return path


def _make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Telemetry"
    sheet.append(["src_ip", "dst_port"])
    sheet.append(["10.0.0.1", 443])
    workbook.save(str(path))
    workbook.close()
    return path


def _artifact_record(source_path: Path, stored_path: Path, sha256: str = "abc123") -> dict[str, object]:
    return {
        "timestamp": "2026-03-08T13:00:00+00:00",
        "source_path": str(source_path),
        "stored_path": str(stored_path),
        "filename": stored_path.name,
        "extension": stored_path.suffix.lower(),
        "mime_type": "application/octet-stream",
        "size_bytes": 42,
        "sha256": sha256,
        "md5": "deadbeef",
        "tags": ["fixture"],
        "risk_level": "medium",
        "reasons": ["test_reason"],
        "extracted_text": "fixture text",
        "extracted_metadata": {"origin": "pytest"},
    }


def test_parse_csv_extracts_shape_and_sample_rows(tmp_path: Path) -> None:
    path = tmp_path / "fixture.csv"
    path.write_text("src_ip,dst_port\n10.0.0.1,443\n10.0.0.2,80\n", encoding="utf-8")

    parsed = parse_csv(path)
    assert parsed["tags"] == ["csv"]
    assert parsed["metadata"]["rows"] == 3
    assert parsed["metadata"]["cols"] == 2
    assert parsed["metadata"]["sample_rows"][1] == ["10.0.0.1", "443"]


def test_parse_json_detects_sensitive_keys(tmp_path: Path) -> None:
    path = tmp_path / "fixture.json"
    path.write_text(json.dumps({"user": "alice", "password": "secret", "cmd": "whoami"}), encoding="utf-8")

    parsed = parse_json(path)
    assert parsed["tags"] == ["json"]
    assert "json_contains_sensitive_or_execution_keys" in parsed["reasons"]
    assert "password" in parsed["metadata"]["suspicious_keys"]


def test_parse_html_detects_suspicious_domains_and_inline_scripts(tmp_path: Path) -> None:
    path = tmp_path / "fixture.html"
    path.write_text(
        """
        <html>
          <head><title>Fixture</title></head>
          <body>
            <script>console.log('1')</script>
            <script>console.log('2')</script>
            <script>console.log('3')</script>
            <script>console.log('4')</script>
            <a href="https://pastebin.com/raw/test">link</a>
          </body>
        </html>
        """,
        encoding="utf-8",
    )

    parsed = parse_html(path)
    assert parsed["tags"] == ["html"]
    assert "html_many_inline_scripts" in parsed["reasons"]
    assert "html_suspicious_domains" in parsed["reasons"]
    assert parsed["metadata"]["inline_scripts"] == 4


def test_parse_python_detects_suspicious_function_calls(tmp_path: Path) -> None:
    path = tmp_path / "fixture.py"
    path.write_text(
        "import os\nimport subprocess\nos.system('whoami')\nsubprocess.run(['ipconfig'])\n",
        encoding="utf-8",
    )

    parsed = parse_python(path)
    assert parsed["tags"] == ["python"]
    assert "python_contains_suspicious_function_calls" in parsed["reasons"]
    assert "os.system" in parsed["metadata"]["suspicious_calls"]


def test_parse_exe_extracts_strings_and_static_indicators(tmp_path: Path) -> None:
    path = tmp_path / "fixture.exe"
    path.write_bytes(b"MZ\x00\x00powershell http://malicious.example cmd.exe")

    parsed = parse_exe(path)
    assert parsed["tags"] == ["exe", "binary"]
    assert "executable_contains_suspicious_strings" in parsed["reasons"]
    assert "powershell" in parsed["metadata"]["suspicious_string_hits"]


def test_parse_zip_detects_dangerous_entries_and_traversal(tmp_path: Path) -> None:
    path = tmp_path / "fixture.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../escape.bat", "echo test")
        archive.writestr("inner/payload.exe", "MZ")
        archive.writestr("nested/archive.zip", "nested")

    parsed = parse_zip(path, max_files=10, max_uncompressed_bytes=1024 * 1024)
    assert parsed["tags"] == ["zip", "archive"]
    assert "zip_contains_dangerous_extensions" in parsed["reasons"]
    assert "zip_contains_nested_archives" in parsed["reasons"]
    assert "zip_path_traversal_indicator" in parsed["reasons"]


def test_parse_docx_pdf_and_xlsx_extract_metadata(tmp_path: Path) -> None:
    docx_path = _make_docx(tmp_path / "fixture.docx")
    pdf_path = _make_pdf(tmp_path / "fixture.pdf")
    xlsx_path = _make_xlsx(tmp_path / "fixture.xlsx")

    docx_parsed = parse_docx(docx_path)
    pdf_parsed = parse_pdf(pdf_path)
    xlsx_parsed = parse_xlsx(xlsx_path)

    assert docx_parsed["tags"] == ["docx"]
    assert docx_parsed["metadata"]["author"] == "pytest"
    assert "deterministic DOCX fixture" in docx_parsed["text"]

    assert pdf_parsed["tags"] == ["pdf"]
    assert pdf_parsed["metadata"]["pages"] == 1
    assert pdf_parsed["metadata"]["/Title"] == "Fixture PDF"

    assert xlsx_parsed["tags"] == ["xlsx"]
    assert xlsx_parsed["metadata"]["sheet_count"] == 1
    assert xlsx_parsed["metadata"]["sheets"][0]["sheet"] == "Telemetry"


def test_artifact_store_migrates_schema_inserts_records_and_writes_jsonl(tmp_path: Path) -> None:
    db_path = tmp_path / "artifacts.db"
    jsonl_path = tmp_path / "artifacts.jsonl"

    with sqlite3.connect(str(db_path)) as conn:
        conn.execute("CREATE TABLE artifacts(id INTEGER PRIMARY KEY, timestamp TEXT, sha256 TEXT)")
        conn.commit()

    source_path = tmp_path / "source.bin"
    stored_path = tmp_path / "stored.bin"
    source_path.write_text("source", encoding="utf-8")
    stored_path.write_text("stored", encoding="utf-8")

    store = ArtifactStore(db_path=db_path, jsonl_path=jsonl_path)
    try:
        assert "stored_path" in store._table_columns("artifacts")
        row_id = store.insert_artifact(_artifact_record(source_path, stored_path))
        found = store.find_by_sha256("abc123")
        assert row_id > 0
        assert found is not None
        assert found["filename"] == "stored.bin"
        assert found["tags"] == ["fixture"]
        assert found["extracted_metadata"]["origin"] == "pytest"
    finally:
        store.close()

    json_payload = json.loads(jsonl_path.read_text(encoding="utf-8").splitlines()[0])
    assert int(json_payload["id"]) == row_id
    assert json_payload["reasons"] == ["test_reason"]


@pytest.mark.integration
def test_run_artifact_scan_moves_duplicates_and_quarantines_high_risk_files(tmp_path: Path) -> None:
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    processed = tmp_path / "processed"
    quarantine = tmp_path / "quarantine"
    db_path = tmp_path / "artifacts.db"
    jsonl_path = tmp_path / "artifacts.jsonl"

    (incoming / "01_safe.json").write_text('{"message": "hello"}', encoding="utf-8")
    (incoming / "02_safe_copy.json").write_text('{"message": "hello"}', encoding="utf-8")
    (incoming / "03_payload.exe").write_bytes(b"MZ powershell http://malicious.example cmd.exe")

    summary = run_artifact_scan(
        path=incoming,
        recursive=False,
        db_path=db_path,
        jsonl_path=jsonl_path,
        processed_dir=processed,
        quarantine_dir=quarantine,
    )

    assert summary == ArtifactScanSummary(
        scanned=3,
        inserted=3,
        duplicates=1,
        quarantined=1,
        processed=2,
        errors=0,
    )
    assert not any(incoming.iterdir())
    assert (processed / "01_safe.json").exists()
    assert (quarantine / "03_payload.exe").exists()

    with sqlite3.connect(str(db_path)) as conn:
        count = conn.execute("SELECT COUNT(*) FROM artifacts").fetchone()[0]
    assert count == 3


def test_generate_artifact_report_summarizes_recorded_results(tmp_path: Path) -> None:
    db_path = tmp_path / "artifacts.db"
    jsonl_path = tmp_path / "artifacts.jsonl"
    store = ArtifactStore(db_path=db_path, jsonl_path=jsonl_path)

    quarantined_path = tmp_path / "quarantine" / "payload.exe"
    quarantined_path.parent.mkdir(parents=True, exist_ok=True)
    quarantined_path.write_text("payload", encoding="utf-8")

    try:
        store.insert_artifact(
            {
                **_artifact_record(tmp_path / "payload.exe", quarantined_path, sha256="sha-risk"),
                "filename": "payload.exe",
                "extension": ".exe",
                "risk_level": "high",
                "reasons": ["executable_contains_suspicious_strings"],
            }
        )
        store.insert_artifact(
            {
                **_artifact_record(tmp_path / "report.json", tmp_path / "processed" / "report.json", sha256="sha-low"),
                "filename": "report.json",
                "extension": ".json",
                "risk_level": "low",
                "reasons": ["json_contains_sensitive_or_execution_keys"],
            }
        )
    finally:
        store.close()

    report_path = generate_artifact_report(db_path, tmp_path / "artifact_report.md")
    report_text = report_path.read_text(encoding="utf-8")

    assert "Total artifacts analyzed: 2" in report_text
    assert "`high` | 1" in report_text
    assert "payload.exe" in report_text
    assert "json_contains_sensitive_or_execution_keys" in report_text


def test_run_artifact_watch_logs_summary_and_stops_cleanly(monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    calls = {"count": 0}

    def fake_scan(*_args: object, **_kwargs: object) -> ArtifactScanSummary:
        calls["count"] += 1
        return ArtifactScanSummary(scanned=1, inserted=1, duplicates=0, quarantined=0, processed=1, errors=0)

    def fake_sleep(_seconds: int) -> None:
        raise KeyboardInterrupt()

    monkeypatch.setattr(intake_module, "run_artifact_scan", fake_scan)
    monkeypatch.setattr(intake_module.time, "sleep", fake_sleep)

    run_artifact_watch(path="incoming", interval_sec=1)
    output = capsys.readouterr().out

    assert calls["count"] == 1
    assert "artifact-watch: scanned=1 inserted=1 duplicates=0 quarantined=0 errors=0" in output
    assert "Artifact watch stopped." in output
